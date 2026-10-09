"""Contact events: mean observed velocity in the 1 s and 2 s before each gripper contact (contact_N.csv -> Contact_Events sheet of final_N.xlsx)."""

import os

import numpy as np
import pandas as pd

from .. import settings
from .step14_velocity import _local_slope_velocity

CSV_COLUMNS = ['Event_ID', 'Contact_Frame', 'Criterion', 'Note']
OUTPUT_COLUMNS = [
    'Event_ID', 'Contact_Frame', 'Criterion', 'Note',
    'Avg_Window_s', 'Avg_Start_Frame', 'Avg_End_Frame',
    'Avg_Vx_cm_s', 'Avg_Vy_cm_s', 'Avg_Vz_cm_s', 'Avg_Speed_cm_s', 'Mean_Speed_cm_s',
    'Obs_Frames_Used', 'Obs_Coverage',
    'Vel_Change_In_Window_cm_s', 'Avg_Speed_k_Spread_cm_s', 'Check_Flags',
]


def contact_csv_path():
    return os.path.join(settings.OUTPUT_DIR, f'contact_{settings.RUN_NUMBER}.csv')


def read_contacts(path):
    if not os.path.exists(path):
        return pd.DataFrame(columns=CSV_COLUMNS)
    for enc in ('utf-8-sig', 'cp949'):
        try:
            df = pd.read_csv(path, encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    for c in CSV_COLUMNS:
        if c not in df.columns:
            df[c] = ''
    df = df.dropna(subset=['Contact_Frame']).copy()
    df['Contact_Frame'] = df['Contact_Frame'].astype(int)
    df[['Criterion', 'Note']] = df[['Criterion', 'Note']].fillna('')
    return df[CSV_COLUMNS]


def write_contacts(path, df):
    df[CSV_COLUMNS].to_csv(path, index=False, encoding='utf-8-sig')


def _config_value(config_df, name, default):
    row = config_df.loc[config_df['Parameter'].astype(str) == name, 'Value']
    try:
        return float(row.iloc[0]) if not row.empty else default
    except (TypeError, ValueError):
        return default


def _window_average(vel, start, end):
    if end < start:
        return np.full(3, np.nan), np.nan, np.nan, 0
    block = vel[start:end + 1]
    ok = np.isfinite(block).all(axis=1)
    if not ok.any():
        return np.full(3, np.nan), np.nan, np.nan, 0
    mean_v = block[ok].mean(axis=0)
    mean_speed = float(np.linalg.norm(block[ok], axis=1).mean())
    return mean_v, float(np.linalg.norm(mean_v)), mean_speed, int(ok.sum())


def _event_row(ev, window_s, frame, time_s, vel_main, vel_by_k, k_main):
    out = {c: np.nan for c in OUTPUT_COLUMNS}
    out.update({'Event_ID': ev['Event_ID'], 'Contact_Frame': ev['Contact_Frame'],
                'Criterion': ev['Criterion'], 'Note': ev['Note'], 'Avg_Window_s': window_s})
    hit = np.flatnonzero(frame == ev['Contact_Frame'])
    if hit.size == 0:
        out['Check_Flags'] = 'FRAME_NOT_FOUND'
        return out
    flags = []
    i_c = int(hit[0])
    start = int(np.searchsorted(time_s, time_s[i_c] - window_s, side='left'))
    end = i_c - k_main
    out['Avg_Start_Frame'] = frame[start]
    out['Avg_End_Frame'] = frame[end] if end >= start else np.nan

    mean_v, avg_speed, mean_speed, n_used = _window_average(vel_main, start, end)
    n_window = max(end - start + 1, 0)
    out.update({'Avg_Vx_cm_s': mean_v[0], 'Avg_Vy_cm_s': mean_v[1], 'Avg_Vz_cm_s': mean_v[2],
                'Avg_Speed_cm_s': avg_speed, 'Mean_Speed_cm_s': mean_speed, 'Obs_Frames_Used': n_used,
                'Obs_Coverage': n_used / n_window if n_window else np.nan})
    if n_used == 0:
        flags.append('NO_DATA')

    if n_window >= 2:
        mid = start + n_window // 2
        first, _, _, n1 = _window_average(vel_main, start, mid - 1)
        second, _, _, n2 = _window_average(vel_main, mid, end)
        if n1 and n2:
            out['Vel_Change_In_Window_cm_s'] = float(np.linalg.norm(second - first))
            if out['Vel_Change_In_Window_cm_s'] > settings.CONTACT_VEL_CHANGE_FLAG_CM_S:
                flags.append('VEL_CHANGE')
        elif n_used:
            flags.append('VEL_CHANGE_NA')

    speeds = []
    for k in settings.CONTACT_SPREAD_HALF_WINDOWS:
        _, s_k, _, n_k = _window_average(vel_by_k[k], start, i_c - k)
        if n_k:
            speeds.append(s_k)
    if len(speeds) >= 2:
        out['Avg_Speed_k_Spread_cm_s'] = float(max(speeds) - min(speeds))
        if out['Avg_Speed_k_Spread_cm_s'] > settings.CONTACT_K_SPREAD_FLAG_CM_S:
            flags.append('K_SPREAD')
    out['Check_Flags'] = ', '.join(flags)
    return out


def compute_contact_events(result, config_df, contacts, windows_s=None):
    windows_s = settings.CONTACT_AVG_WINDOWS_S if windows_s is None else windows_s
    k_main = int(_config_value(config_df, 'VEL_HALF_WINDOW_frames', settings.VEL_HALF_WINDOW))
    max_ratio = _config_value(config_df, 'VEL_MAX_NOISE_RATIO', settings.VEL_MAX_NOISE_RATIO)
    max_offset = _config_value(config_df, 'VEL_MAX_CENTER_OFFSET_frames', settings.VEL_MAX_CENTER_OFFSET_FRAMES)

    frame = result['Frame'].to_numpy()
    time_s = result['Time_s'].to_numpy(dtype=float)
    coords = result[['X_cm', 'Y_cm', 'Z_cm']].to_numpy(dtype=float)
    floor = result['Z_Floor_Applied'].fillna(False).astype(bool).to_numpy()
    state = result['State'].astype(str).to_numpy()
    finite = np.isfinite(coords).all(axis=1)
    usable_obs = finite & ~floor & (result['Coord_Type'].astype(str).to_numpy() == 'STEREO')
    output_mask = finite & ~floor & np.isin(state, ['DETECTED', 'RECOVERED'])

    vel_by_k = {k: _local_slope_velocity(time_s, coords, usable_obs, output_mask, k,
                                         max_ratio=max_ratio, max_center_offset=max_offset)[0]
                for k in sorted(set(settings.CONTACT_SPREAD_HALF_WINDOWS) | {k_main})}
    vel_main = result[['Vx_obs_cm_s', 'Vy_obs_cm_s', 'Vz_obs_cm_s']].to_numpy(dtype=float)
    if not np.allclose(vel_by_k[k_main], vel_main, equal_nan=True, atol=1e-9):
        print('WARNING: recomputed observed velocity differs from V*_obs_cm_s in the final Excel. Check the Configuration sheet.')

    rows = [_event_row(ev, w, frame, time_s, vel_main, vel_by_k, k_main)
            for _, ev in contacts.iterrows() for w in windows_s]
    return pd.DataFrame(rows, columns=OUTPUT_COLUMNS)


def _format_sheet(ws):
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    ws.freeze_panes = 'A2'
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col_idx, cells in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        header = str(cells[0].value or '')
        letter = get_column_letter(col_idx)
        if header in ('Criterion', 'Note'):
            ws.column_dimensions[letter].width = 40
            for c in cells[1:]:
                c.alignment = Alignment(wrap_text=True, vertical='top')
        else:
            longest = max([len(header)] + [len(str(c.value)) for c in cells[1:] if c.value is not None])
            ws.column_dimensions[letter].width = min(max(longest, 6) + 2, 40)


def run_contact_events(final_path=None, csv_path=None):
    final_path = final_path or settings.FINAL_EXCEL
    csv_path = csv_path or contact_csv_path()
    contacts = read_contacts(csv_path)
    if contacts.empty:
        print(f'No contact events in {csv_path}. Enter them in the form and click Save.')
        return None
    result = pd.read_excel(final_path, sheet_name='Result')
    config_df = pd.read_excel(final_path, sheet_name='Configuration')
    events = compute_contact_events(result, config_df, contacts)

    with pd.ExcelWriter(final_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        events.to_excel(writer, sheet_name='Contact_Events', index=False)
        wb = writer.book
        ws = wb['Contact_Events']
        _format_sheet(ws)
        target = wb.sheetnames.index('Result') + 1 if 'Result' in wb.sheetnames else 0
        wb.move_sheet(ws, offset=target - wb.sheetnames.index('Contact_Events'))

    print(f'Contact_Events saved: {final_path} ({len(contacts)} events x {len(settings.CONTACT_AVG_WINDOWS_S)} windows)')
    flagged = events[events['Check_Flags'].astype(str).str.len() > 0]
    if len(flagged):
        print('Recheck these events in the video:')
        print(flagged[['Event_ID', 'Contact_Frame', 'Avg_Window_s', 'Check_Flags']].to_string(index=False))
    return events
